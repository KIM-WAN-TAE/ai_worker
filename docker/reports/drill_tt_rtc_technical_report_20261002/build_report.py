#!/usr/bin/env python3
"""Build the template-based Drill TT-RTC DOCX and its diagram assets."""

from __future__ import annotations

import math
import re
import shutil
import sys
import textwrap
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from PIL import Image, ImageDraw, ImageFont, ImageOps


ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
ASSETS = OUT / "assets"
TEMPLATE_DIR = OUT / "template_reference"
SOURCE = ROOT / "docs/DRILL_TT_RTC_TECHNICAL_REPORT_2026-10-02.md"
PHOTO_SOURCE = ROOT / "docs/assets/drill_tt_rtc/drill_insertion_episode_000000_frame_000184.jpg"
DOTX_SOURCE = TEMPLATE_DIR / "Dara_Technical_Project_Report_Template.dotx"
OUTPUT_DOCX = OUT / "DRILL_TT_RTC_TECHNICAL_REPORT_2026-10-02.docx"
WORKING_TEMPLATE = TEMPLATE_DIR / "Dara_Technical_Project_Report_Template.docx"

FONT_PATH = "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"
BOLD_FONT_PATH = "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc"
LANCZOS = getattr(getattr(Image, "Resampling", Image), "LANCZOS")

sys.path.insert(0, "/tmp/drill_rlt_docx_deps")
from docx import Document  # noqa: E402
from docx.enum.section import WD_SECTION  # noqa: E402
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Inches, Pt, RGBColor  # noqa: E402


COLORS = {
    "ink": "172B3A",
    "navy": "173E5E",
    "blue": "2F7399",
    "sky": "EAF4F8",
    "teal": "287765",
    "mint": "EAF5F1",
    "orange": "C06B32",
    "peach": "FFF1E7",
    "red": "A23B3B",
    "gray": "5F6B73",
    "light": "F4F6F7",
    "line": "A9B3BA",
    "white": "FFFFFF",
}


def pil_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    path = BOLD_FONT_PATH if bold else FONT_PATH
    return ImageFont.truetype(path, size=size)


class Diagram:
    def __init__(self, title: str, width: int = 2000, height: int = 1250):
        self.width = width
        self.height = height
        self.image = Image.new("RGB", (width, height), "white")
        self.draw = ImageDraw.Draw(self.image)
        self.draw.text(
            (width // 2, 45),
            title,
            anchor="ma",
            fill="#173E5E",
            font=pil_font(52, True),
        )

    def text(self, xy, value, size=29, color="#172B3A", bold=False, anchor="mm"):
        self.draw.multiline_text(
            xy,
            value,
            anchor=anchor,
            align="center",
            spacing=9,
            fill=color,
            font=pil_font(size, bold),
        )

    def box(self, x, y, w, h, title, detail="", fill="#EAF4F8", outline="#2F7399"):
        self.draw.rounded_rectangle(
            (x, y, x + w, y + h), radius=23, fill=fill, outline=outline, width=5
        )
        self.text((x + w / 2, y + h * (0.38 if detail else 0.5)), title, 35, bold=True)
        if detail:
            self.text((x + w / 2, y + h * 0.73), detail, 25, color="#5F6B73")
        return (x, y, w, h)

    def arrow(self, start, end, color="#7B8891", label=None, offset=(0, -24), width=7):
        x1, y1 = start
        x2, y2 = end
        self.draw.line((x1, y1, x2, y2), fill=color, width=width)
        angle = math.atan2(y2 - y1, x2 - x1)
        head = 23
        spread = 0.55
        self.draw.polygon(
            [
                (x2, y2),
                (x2 - head * math.cos(angle - spread), y2 - head * math.sin(angle - spread)),
                (x2 - head * math.cos(angle + spread), y2 - head * math.sin(angle + spread)),
            ],
            fill=color,
        )
        if label:
            self.text(
                ((x1 + x2) / 2 + offset[0], (y1 + y2) / 2 + offset[1]),
                label,
                24,
                color="#5F6B73",
            )

    def save(self, name: str) -> Path:
        path = ASSETS / name
        self.image.save(path, dpi=(220, 220), optimize=True)
        return path


def make_diagrams() -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)
    shutil.copy2(PHOTO_SOURCE, ASSETS / PHOTO_SOURCE.name)

    d = Diagram("Drill TT-RTC 전체 구조", height=1420)
    y = 160
    boxes = [
        d.box(80, y, 300, 160, "시연 데이터", "3 cameras · 16D state\n32×16 action · language", "#FFF1E7", "#C06B32"),
        d.box(470, y, 300, 160, "TT-RTC 학습", "d=0…6 clean prefix\npostfix-only flow loss"),
        d.box(860, y, 300, 160, "Screw_100k", "step 100,000\nmanifest + processor"),
        d.box(1250, y, 300, 160, "추론 엔진", "4-step flow\nprefix 매 step clamp", "#EAF5F1", "#287765"),
        d.box(1640, y, 280, 160, "32×16 출력", "prefix 6\npostfix 26", "#EAF5F1", "#287765"),
    ]
    for left, right in zip(boxes, boxes[1:]):
        d.arrow((left[0] + left[2], y + 80), (right[0], y + 80))

    d.box(100, 520, 350, 180, "실시간 관측", "head · left wrist · right wrist\narm_left/right · instruction", "#FFF1E7", "#C06B32")
    d.box(600, 520, 360, 180, "Main control loop", "queue≤6이면 prefix capture\nrequest · 검증 · atomic commit")
    d.box(1110, 520, 350, 180, "TTActionTimeline", "15 Hz source waypoint\ncontrol clock 보간", "#EAF5F1", "#287765")
    d.box(1610, 520, 300, 180, "SG2 로봇", "양팔 Drill 실행\ncommand trace", "#EAF5F1", "#287765")
    d.arrow((450, 610), (600, 610))
    d.arrow((960, 610), (1110, 610))
    d.arrow((1460, 610), (1610, 610))
    d.arrow((780, 520), (1380, 320), label="prefix-conditioned request", offset=(0, -30))
    d.arrow((1780, 700), (1780, 950), color="#C06B32", label="실행 중 남은 source action", offset=(-170, 0))
    d.arrow((1780, 950), (790, 700), color="#C06B32", label="exact committed prefix feedback", offset=(0, 32))
    d.box(210, 1030, 1580, 210, "명시적 계약 경계", "logical 32×16  |  model 40×132  |  source 15 Hz  |  d∈[0,6]  |  clean endpoint=1000\nshape · finite · prefix equality · queue suffix · deadline · sensor freshness 검증", "#F4F6F7", "#7B8891")
    d.text((1000, 1335), "현재 활성 경로: VLA policy + TT-RTC  |  RL Critic · reward optimizer 미활성", 29, "#A23B3B", True)
    d.save("diagram_01_overall.png")

    d = Diagram("TT-RTC 실행 시퀀스: bootstrap과 정상 refill", height=1250)
    x_positions = [100, 465, 830, 1195, 1560]
    names = ["Main runtime", "Source queue", "Inference engine", "GR00T TT-RTC", "로봇"]
    for x, name in zip(x_positions, names):
        d.text((x + 130, 155), name, 29, bold=True)
        d.draw.line((x + 130, 205, x + 130, 1130), fill="#CDD4D9", width=4)
    def message(y, src, dst, text_value, color="#2F7399"):
        d.arrow((x_positions[src] + 130, y), (x_positions[dst] + 130, y), color=color, label=text_value, offset=(0, -30), width=5)
    message(270, 0, 2, "① bootstrap: d=0, prefix=[]")
    message(350, 2, 3, "② 32-step 생성")
    message(430, 3, 1, "③ 검증 후 32개 enqueue", "#287765")
    message(520, 1, 4, "④ 15 Hz source → control tick", "#C06B32")
    message(650, 1, 0, "⑤ 남은 source action≤6")
    message(730, 0, 2, "⑥ d=6, exact 6×16 prefix")
    message(810, 2, 3, "⑦ prefix-conditioned 32-step")
    message(890, 3, 0, "⑧ prefix equality · deadline")
    message(970, 0, 1, "⑨ 기존 6 유지 + postfix 26 append", "#287765")
    message(1050, 1, 4, "⑩ 연속 실행", "#C06B32")
    d.text((1000, 1195), "오류 시 새 postfix를 폐기하고 failure를 latch한 뒤 명시적 RESUME을 요구", 27, "#A23B3B", True)
    d.save("diagram_02_runtime_sequence.png")

    d = Diagram("15 Hz source clock과 control clock의 분리", height=1030)
    left = 180
    right = 1840
    y_source = 320
    y_control = 670
    d.draw.line((left, y_source, right, y_source), fill="#173E5E", width=8)
    for i in range(10):
        x = left + i * (right - left) / 9
        d.draw.ellipse((x - 15, y_source - 15, x + 15, y_source + 15), fill="#2F7399")
        d.text((x, y_source - 65), f"q{i}", 24)
    d.text((100, y_source), "source\n15 Hz", 28, bold=True)
    d.draw.line((left, y_control, right, y_control), fill="#C06B32", width=5)
    for i in range(61):
        x = left + i * (right - left) / 60
        d.draw.line((x, y_control - 13, x, y_control + 13), fill="#C06B32", width=3)
    d.text((100, y_control), "control\n예: 100 Hz", 28, bold=True)
    d.arrow((left + 6 * (right - left) / 9, y_source + 30), (left + 6 * (right - left) / 9, 520), color="#287765", label="peek_actions(): exact source prefix", offset=(230, 0))
    d.arrow((left + 7 * (right - left) / 9, y_source + 30), (left + 7 * (right - left) / 9, 520), color="#287765")
    d.text((1000, 805), "pop_action()만 인접 source waypoint를 보간한다.\n보간된 control command를 다음 TT-RTC prefix로 재사용하지 않는다.", 31, "#173E5E", True)
    d.save("diagram_03_timeline.png")

    d = Diagram("TT-RTC 안전 상태 전이", height=1180)
    nodes = {
        "load": d.box(120, 230, 300, 150, "LOAD", "manifest·model·processor\n계약 검증", "#EAF4F8", "#2F7399"),
        "boot": d.box(600, 230, 300, 150, "BOOTSTRAP", "queue empty\nd=0", "#FFF1E7", "#C06B32"),
        "run": d.box(1080, 230, 300, 150, "RUN", "source action 실행\nqueue monitoring", "#EAF5F1", "#287765"),
        "refill": d.box(1560, 230, 300, 150, "REFILL", "queue≤6\nd=6 request", "#EAF5F1", "#287765"),
        "pause": d.box(1080, 730, 430, 180, "LATCHED PAUSE", "새 continuation 폐기\n명시적 START/RESUME 필요", "#FCEBEC", "#A23B3B"),
        "stop": d.box(300, 730, 350, 180, "STOP/UNLOAD", "queue·session 정리\nold response 격리", "#F4F6F7", "#7B8891"),
    }
    d.arrow((420, 305), (600, 305), label="성공")
    d.arrow((900, 305), (1080, 305), label="32개 commit")
    d.arrow((1380, 305), (1560, 305), label="남은 6개")
    d.arrow((1710, 380), (1240, 230), color="#287765", label="prefix 6 + postfix 26", offset=(0, 38))
    d.arrow((1710, 380), (1420, 730), color="#A23B3B", label="검증·deadline 실패", offset=(70, 0))
    d.arrow((1230, 380), (1230, 730), color="#A23B3B", label="sensor stale · queue exhausted", offset=(170, 0))
    d.arrow((1080, 820), (650, 820), color="#7B8891", label="명시적 stop")
    d.arrow((1295, 730), (750, 380), color="#C06B32", label="원인 제거 후 explicit resume", offset=(0, 28))
    d.save("diagram_04_safety.png")

    d = Diagram("VLA · TT-RTC · 강화학습 검증 피라미드", height=1450)
    levels = [
        ("실로봇 Drill A/B/C", "성공률 · 축 이탈 · jerk · intervention · safety cost", 420, "#FFF1E7", "#C06B32"),
        ("runtime 통합", "bootstrap · refill · deadline · queue suffix · safety shield", 640, "#EAF5F1", "#287765"),
        ("simulation RL", "multi-seed · recovery · reward ablation · constraint", 810, "#EAF4F8", "#2F7399"),
        ("Critic · 정책 검증", "OPE · calibration · Q 과대평가 · PPO ratio · KL", 980, "#EAF4F8", "#2F7399"),
        ("loss · 계약 단위", "reward · GAE · postfix mask · denominator · padding", 1200, "#F4F6F7", "#7B8891"),
    ]
    center = 1000
    y = 140
    for title, detail, width, fill, outline in levels:
        x = center - width / 2
        d.box(x, y, width, 160, title, detail, fill, outline)
        y += 205
    d.text((1000, 1320), "아래 계층을 통과하지 않은 결과로 위 계층의 배포 적합성을 주장하지 않는다.", 28, "#A23B3B", True)
    d.save("diagram_05_validation.png")

    d = Diagram("VLA · 강화학습 · TT-RTC 결합 구조", height=1450)
    d.box(90, 170, 330, 175, "경험 데이터", "시연 · 실패 · simulation\n실로봇 rollout", "#FFF1E7", "#C06B32")
    d.box(540, 170, 330, 175, "Reward · Cost", "성공 · 정렬 · 삽입\n충돌 · 힘 · jerk", "#FFF1E7", "#C06B32")
    d.box(990, 170, 330, 175, "Replay / Rollout", "(s, a, r, c, s′)\ntrajectory provenance", "#F4F6F7", "#7B8891")
    d.box(1440, 170, 390, 175, "Critic · Value", "Q(s,a) · V(s)\nadvantage · constraint", "#EAF4F8", "#2F7399")
    d.arrow((420, 258), (540, 258))
    d.arrow((870, 258), (990, 258))
    d.arrow((1320, 258), (1440, 258))

    d.box(220, 570, 430, 195, "VLA Actor 업데이트", "PPO / SAC / IQL · AWAC\nBC·KL·TT-RTC 보조 loss", "#EAF4F8", "#2F7399")
    d.box(785, 570, 430, 195, "RL 검증 체크포인트", "offline OPE · simulation\nsafety gate · A/B", "#EAF5F1", "#287765")
    d.box(1350, 570, 430, 195, "TT-RTC Runtime", "committed prefix 유지\npostfix 실시간 생성", "#EAF5F1", "#287765")
    d.arrow((1635, 345), (520, 570), label="advantage / TD target", offset=(0, -26))
    d.arrow((650, 668), (785, 668))
    d.arrow((1215, 668), (1350, 668))

    d.box(1320, 975, 490, 205, "Drill 로봇 + Safety Shield", "workspace · 속도 · 충돌 · 힘 제한\nemergency stop · human supervision", "#FCEBEC", "#A23B3B")
    d.box(190, 975, 470, 205, "평가 및 재수집", "성공률 · boundary jerk\nlateral deviation · intervention", "#FFF1E7", "#C06B32")
    d.arrow((1565, 765), (1565, 975), color="#287765", label="배포 actor 실행", offset=(150, 0))
    d.arrow((1320, 1078), (660, 1078), color="#C06B32", label="trajectory · reward · cost", offset=(0, -30))
    d.arrow((425, 975), (255, 345), color="#C06B32", label="검증된 데이터만 재학습", offset=(-35, 0))
    d.text((1000, 1320), "현재 상태: TT-RTC는 활성 · Critic/Reward/RL update는 미활성   |   본 구조: 단계적 RL 확장 설계안", 28, "#A23B3B", True)
    d.save("diagram_06_rl_extension.png")


def convert_template_package() -> None:
    """Instantiate the downloaded DOTX as DOCX while preserving its styles."""
    TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)
    if not DOTX_SOURCE.is_file():
        raise FileNotFoundError(f"technical report template is missing: {DOTX_SOURCE}")
    with ZipFile(DOTX_SOURCE, "r") as source, ZipFile(
        WORKING_TEMPLATE, "w", compression=ZIP_DEFLATED
    ) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "[Content_Types].xml":
                data = data.replace(
                    b"application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml",
                    b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
                )
            target.writestr(info, data)


def clear_document(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_cell_shading(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_width(cell, width_twips: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, value: str, bold: bool = False, color: str | None = None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(value)
    run.bold = bold
    run.font.name = "Noto Sans CJK KR"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK KR")
    run.font.size = Pt(8.3)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def remove_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "nil")


def add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("CYCLO-TR-TTRTC-DRILL-2026-001  ·  ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(COLORS["gray"])
    add_field(paragraph, " PAGE ", "1")


def configure_styles(doc: Document) -> None:
    style_specs = {
        "Normal": ("Noto Serif CJK KR", 9.3, False, COLORS["ink"]),
        "Body Text": ("Noto Serif CJK KR", 9.3, False, COLORS["ink"]),
        "Title": ("Noto Sans CJK KR", 27, True, COLORS["navy"]),
        "Subtitle": ("Noto Sans CJK KR", 12, False, COLORS["gray"]),
        "Heading 1": ("Noto Sans CJK KR", 18, True, COLORS["navy"]),
        "Heading 2": ("Noto Sans CJK KR", 13, True, COLORS["blue"]),
        "Heading 3": ("Noto Sans CJK KR", 10.5, True, COLORS["ink"]),
        "Caption": ("Noto Sans CJK KR", 8.5, False, COLORS["gray"]),
        "Code": ("Noto Sans Mono CJK KR", 7.7, False, COLORS["ink"]),
        "Quote": ("Noto Serif CJK KR", 9.2, False, COLORS["navy"]),
        "Bibliography": ("Noto Serif CJK KR", 8.5, False, COLORS["ink"]),
    }
    for name, (font_name, size, bold, color) in style_specs.items():
        if name not in doc.styles:
            continue
        style = doc.styles[name]
        style.font.name = font_name
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
        style._element.rPr.rFonts.set(qn("w:ascii"), font_name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), font_name)
    # The downloaded template carries automatic outline numbering.  This
    # report keeps the explicit Korean chapter numbers in its source so that
    # unnumbered front matter and numbered main matter coexist cleanly.
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        p_pr = style._element.get_or_add_pPr()
        num_pr = p_pr.find(qn("w:numPr"))
        if num_pr is not None:
            p_pr.remove(num_pr)
    normal = doc.styles["Normal"].paragraph_format
    normal.line_spacing = 1.32
    normal.space_after = Pt(5.5)
    normal.widow_control = True
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(18)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(8)
    doc.styles["Heading 1"].paragraph_format.keep_with_next = True
    for name in ("Heading 2", "Heading 3"):
        doc.styles[name].paragraph_format.space_before = Pt(10)
        doc.styles[name].paragraph_format.space_after = Pt(4)
        doc.styles[name].paragraph_format.keep_with_next = True


def add_inline_text(paragraph, text: str) -> None:
    token_re = re.compile(r"(\*\*.*?\*\*|`.*?`|\$[^$]+\$|https?://\S+)")
    cursor = 0
    for match in token_re.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor : match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Noto Sans Mono CJK KR"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans Mono CJK KR")
            run.font.size = Pt(8.2)
            run.font.color.rgb = RGBColor.from_string(COLORS["blue"])
        elif token.startswith("$"):
            formula, _ = equation_plain(token)
            run = paragraph.add_run(formula)
            run.font.name = "Cambria Math"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Serif CJK KR")
            run.font.size = Pt(9.2)
        else:
            run = paragraph.add_run(token.rstrip(".,"))
            run.font.color.rgb = RGBColor.from_string("0563C1")
            run.underline = True
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def add_caption(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = False
    add_inline_text(paragraph, text)


def add_table(doc: Document, rows: list[list[str]], caption: str) -> None:
    add_caption(doc, caption)
    columns = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for r_index, row in enumerate(rows):
        for c_index in range(columns):
            value = row[c_index] if c_index < len(row) else ""
            set_cell_text(
                table.cell(r_index, c_index),
                value,
                bold=(r_index == 0),
                color=(COLORS["white"] if r_index == 0 else None),
            )
            if r_index == 0:
                set_cell_shading(table.cell(r_index, c_index), COLORS["navy"])
            elif r_index % 2 == 0:
                set_cell_shading(table.cell(r_index, c_index), "F3F7F9")
    set_repeat_table_header(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def balanced_group(value: str, start: int) -> tuple[str, int] | None:
    if start >= len(value) or value[start] != "{":
        return None
    depth = 0
    for index in range(start, len(value)):
        if value[index] == "{":
            depth += 1
        elif value[index] == "}":
            depth -= 1
            if depth == 0:
                return value[start + 1 : index], index + 1
    return None


def replace_command_groups(value: str, command: str) -> str:
    marker = f"\\{command}"
    while marker in value:
        pos = value.find(marker)
        group = balanced_group(value, pos + len(marker))
        if group is None:
            value = value.replace(marker, "", 1)
            continue
        content, end = group
        value = value[:pos] + content + value[end:]
    return value


def replace_fractions(value: str) -> str:
    marker = "\\frac"
    while marker in value:
        pos = value.rfind(marker)
        first = balanced_group(value, pos + len(marker))
        if first is None:
            value = value[:pos] + value[pos + len(marker) :]
            continue
        numerator, next_pos = first
        second = balanced_group(value, next_pos)
        if second is None:
            value = value[:pos] + numerator + value[next_pos:]
            continue
        denominator, end = second
        value = value[:pos] + f"({numerator})/({denominator})" + value[end:]
    return value


def equation_plain(latex: str) -> tuple[str, str]:
    value = latex.strip()
    if value.startswith("$$"):
        value = value[2:]
    elif value.startswith("$"):
        value = value[1:]
    if value.endswith("$$"):
        value = value[:-2]
    elif value.endswith("$"):
        value = value[:-1]
    value = value.strip()
    number = ""
    match = re.search(r"\\tag\{(\d+)\}", value)
    if match:
        number = match.group(1)
        value = value[: match.start()] + value[match.end() :]
    else:
        match = re.search(r"\\quad\s*\((\d+)\)\s*$", value)
    if match:
        if not number:
            number = match.group(1)
            value = value[: match.start()].strip()
    value = value.replace("\\dfrac", "\\frac")
    value = re.sub(r"\\\\(?:\[[^\]]*\])?", " @@EQUATION_LINE@@ ", value)
    value = value.replace("\\begin{aligned}", "").replace("\\end{aligned}", "")
    value = value.replace("\\leftarrow", "←")
    value = value.replace("\\left\\|", "‖").replace("\\right\\|", "‖")
    value = value.replace("\\left", "").replace("\\right", "")
    value = value.replace("\\qquad", "   ").replace("\\quad", "   ")
    value = value.replace("\\;", " ").replace("\\,", " ").replace("\\!", "")
    value = value.replace("\\{", "⦃").replace("\\}", "⦄")
    value = value.replace("\\begin{cases}", "⎧ ").replace("\\end{cases}", "")
    value = value.replace("\\mathcal{L}", "ℒ").replace("\\mathcal{N}", "N")
    value = value.replace("\\mathbf{1}", "1")
    value = value.replace("\\min", "min")
    value = replace_fractions(value)
    for command in (
        "operatorname",
        "widehat",
        "widetilde",
        "boldsymbol",
        "mathbb",
        "mathrm",
        "mathbf",
        "text",
        "mathcal",
    ):
        value = replace_command_groups(value, command)
    replacements = {
        "\\longrightarrow": "→",
        "\\sum": "Σ",
        "\\sim": "∼",
        "\\le": "≤",
        "\\ge": "≥",
        "\\in": "∈",
        "\\land": "∧",
        "\\Vert": "‖",
        "\\oplus": "⊕",
        "\\forall": "∀",
        "\\Pr": "P",
        "\\ldots": "…",
        "\\varepsilon": "ε",
        "\\hat": "",
        "\\tilde": "",
        "\\theta": "θ",
        "\\tau": "τ",
        "\\pi": "π",
        "\\ell": "ℓ",
        "\\alpha": "α",
        "\\beta": "β",
        "\\gamma": "γ",
        "\\delta": "δ",
        "\\lambda": "λ",
        "\\omega": "ω",
        "\\rho": "ρ",
        "\\phi": "φ",
        "\\psi": "ψ",
        "\\Phi": "Φ",
        "\\mid": "|",
        "\\cdot": "·",
        "\\epsilon": "ε",
        "\\Delta": "Δ",
        "\\partial": "∂",
        "\\mathbb{R}": "ℝ",
        "\\mathcal": "",
        "\\mathrm": "",
        "\\displaystyle": "",
        "\\#": "#",
        "\\ ": " ",
        "\\textstyle": "",
    }
    for old, new in replacements.items():
        value = value.replace(old, new)
    value = value.replace("&", "")
    value = value.replace("^{", "^(").replace("_{", "_(")
    value = value.replace("}", ")").replace("{", "(")
    value = value.replace("\\", "")
    value = re.sub(r"\s+", " ", value).strip()
    value = value.replace(" @@EQUATION_LINE@@ ", "\n")
    return value, number


def add_equation(doc: Document, latex: str) -> None:
    formula, number = equation_plain(latex)
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    remove_table_borders(table)
    set_cell_width(table.cell(0, 0), 8500)
    set_cell_width(table.cell(0, 1), 700)
    p = table.cell(0, 0).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(formula)
    run.font.name = "Cambria Math"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Serif CJK KR")
    longest_line = max((len(line) for line in formula.splitlines()), default=0)
    run.font.size = Pt(9.2 if longest_line < 120 else 7.7)
    pn = table.cell(0, 1).paragraphs[0]
    pn.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    number_run = pn.add_run(f"({number})" if number else "")
    number_run.font.name = "Cambria Math"
    number_run.font.size = Pt(9)


def resolve_image(path_text: str) -> Path:
    path = (SOURCE.parent / path_text).resolve()
    if not path.is_file():
        raise FileNotFoundError(path)
    return path


def parse_table(lines: list[str], index: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    while index < len(lines) and lines[index].strip().startswith("|"):
        row = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        rows.append(row)
        index += 1
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in rows[1]):
        rows.pop(1)
    return rows, index


def scan_lists(
    lines: list[str], *, table_start: int = 0
) -> tuple[list[str], list[str], list[str]]:
    figures: list[str] = []
    tables: list[str] = []
    equations: list[str] = []
    heading = ""
    table_no = table_start
    for index, line in enumerate(lines):
        if line.startswith("# ") or line.startswith("## ") or line.startswith("### "):
            heading = line.lstrip("# ").strip()
        if line.startswith("그림 "):
            figures.append(line.strip())
        if line.strip().startswith("|") and (
            index == 0 or not lines[index - 1].strip().startswith("|")
        ):
            table_no += 1
            tables.append(f"표 {table_no} — {heading or '기술 데이터'}")
        tag_match = re.search(r"\\tag\{(\d+)\}", line)
        if tag_match:
            number = tag_match.group(1)
            equations.append(f"수식 {number} — {heading}")
        elif (
            line.strip() != "$$"
            and line.strip().startswith("$$")
            and line.strip().endswith("$$")
        ):
            _, number = equation_plain(line.strip())
            equations.append(f"수식 {number} — {heading}" if number else f"수식 — {heading}")
    return figures, tables, equations


def add_cover(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.different_first_page_header_footer = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("CYCLO INTELLIGENCE  /  TECHNICAL REPORT")
    r.font.name = "Noto Sans CJK KR"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK KR")
    r.font.size = Pt(9)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(COLORS["blue"])
    p.paragraph_format.space_after = Pt(35)

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(8)
    title.add_run("Drill 삽입 작업을 위한\nVLA · TT-RTC · 강화학습 기술보고서")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("Training-Time Real-Time Chunking & Reinforcement Learning\n알고리즘 · Reward · Loss · 실행 구조 · 검증 계획")
    subtitle.paragraph_format.space_after = Pt(20)

    image = Image.open(PHOTO_SOURCE).convert("RGB")
    cover = ImageOps.fit(image, (1680, 850), method=LANCZOS, centering=(0.5, 0.5))
    overlay = Image.new("RGBA", cover.size, (23, 62, 94, 0))
    draw = ImageDraw.Draw(overlay)
    draw.rectangle((0, 610, 1680, 850), fill=(23, 62, 94, 205))
    draw.text((55, 665), "DRILL INTO SCREW", fill="white", font=pil_font(52, True))
    draw.text((55, 755), "Actual scene · episode 0 · head camera · frame 184", fill="#DCEAF2", font=pil_font(28))
    cover = Image.alpha_composite(cover.convert("RGBA"), overlay).convert("RGB")
    cover_path = ASSETS / "drill_insertion_cover.jpg"
    cover.save(cover_path, quality=94, dpi=(220, 220))
    photo_p = doc.add_paragraph()
    photo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    photo_p.add_run().add_picture(str(cover_path), width=Cm(17.0))

    doc.add_paragraph()
    control = doc.add_table(rows=5, cols=2)
    control.style = "Table Grid"
    control.alignment = WD_TABLE_ALIGNMENT.CENTER
    values = [
        ("문서 번호", "CYCLO-TR-TTRTC-DRILL-2026-001"),
        ("기준일", "2026-10-02"),
        ("문서 상태", "Final — 현재 TT-RTC 검증 및 강화학습 확장 설계 통합본"),
        ("적용 시스템", "CYCLO Intelligence / GR00T N1.7 / SG2"),
        ("핵심 범위", "현재 TT-RTC 분석 · 강화학습 핵심 원리 및 결합 설계"),
    ]
    for i, (key, value) in enumerate(values):
        set_cell_text(control.cell(i, 0), key, True, COLORS["white"])
        set_cell_shading(control.cell(i, 0), COLORS["navy"])
        set_cell_text(control.cell(i, 1), value)
    doc.add_page_break()


def add_front_matter(doc: Document, lines: list[str]) -> None:
    heading = doc.add_paragraph("문서 관리", style="Heading 1")
    heading.paragraph_format.page_break_before = False
    rows = [
        ["항목", "내용"],
        ["작성 목적", "Drill VLA·TT-RTC·강화학습 알고리즘과 검증 구조의 단일 기준 문서"],
        ["작성 근거", "체크포인트 manifest/config, runtime source, tests, 공개 논문·작성 지침"],
        ["Template", "Dara’s Technical Project Report Template (CC0-1.0)"],
        ["검토 필요", "실로봇 성공률·접촉 안정성·배포 적합성은 별도 시험으로 확정"],
        ["변경 통제", "H/D/d/rate/loss/processor/dataset 변경 시 재검증"],
    ]
    add_table(doc, rows, "표 1 — 문서 관리 정보")
    q = doc.add_paragraph(style="Quote")
    q.add_run("범위 선언: 현재 Drill 활성 경로는 base GR00T + TT-RTC이다. 강화학습 장은 PPO·SAC·IQL/AWAC의 원리와 향후 결합 설계를 상세히 다루지만, 현재 체크포인트에 Critic·Reward optimizer가 구현되어 있다는 뜻은 아니다.")
    doc.add_page_break()

    doc.add_paragraph("초록", style="Heading 1")
    abstract = (
        "본 보고서는 양팔 로봇의 Drill 삽입 작업을 대상으로 VLA, TT-RTC, 강화학습의 역할을 하나의 "
        "구조로 설명한다. 현재 체크포인트는 15 Hz에서 32×16 논리 action chunk를 생성하고, "
        "0~6 step의 clean prefix를 조건으로 postfix-only flow-matching loss를 계산한다. 여기에 "
        "강화학습을 결합할 때 필요한 MDP, reward와 cost, Actor–Critic, PPO, SAC, IQL/AWAC, "
        "advantage-weighted TT-RTC loss를 수식과 함께 제시한다. RL은 성공률·정렬·삽입 깊이 같은 "
        "작업 결과를 직접 최적화하고 실패 복구를 학습할 수 있지만, reward hacking과 위험한 탐색을 "
        "통제해야 한다. 따라서 현재 활성 TT-RTC와 향후 RL 설계안을 분리해 표시하고, offline 학습, "
        "simulation 검증, 제한된 실로봇 fine-tuning, safety gate의 단계로 배포 절차를 정의한다."
    )
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0.7)
    add_inline_text(p, abstract)
    p = doc.add_paragraph()
    p.add_run("키워드: ").bold = True
    p.add_run("TT-RTC, VLA, reinforcement learning, PPO, SAC, offline RL, flow matching, Drill insertion")
    doc.add_page_break()

    doc.add_paragraph("목차", style="Heading 1")
    toc = doc.add_paragraph()
    add_field(toc, ' TOC \\o "1-3" \\h \\z \\u ', "문서 열기 후 목차 업데이트")
    doc.add_page_break()

    body_start = next(i for i, line in enumerate(lines) if line.startswith("# 1. 서론"))
    figures, body_tables, equations = scan_lists(lines[body_start:], table_start=2)
    tables = [
        "표 1 — 문서 관리 정보",
        "표 2 — 용어와 표기",
        *body_tables,
    ]
    doc.add_paragraph("그림·표·수식 목록", style="Heading 1")
    for title, items in (("그림", figures), ("표", tables), ("수식", equations)):
        doc.add_paragraph(title, style="Heading 2")
        for item in items:
            doc.add_paragraph(item, style="Caption")
    doc.add_page_break()

    doc.add_paragraph("용어와 표기", style="Heading 1")
    glossary = [
        ["용어", "정의"],
        ["VLA", "영상·언어·로봇 상태로부터 action chunk를 생성하는 정책"],
        ["TT-RTC", "clean committed prefix를 조건으로 postfix를 생성하도록 학습한 실시간 청크 방식"],
        ["prefix", "추론 중에도 큐에 남아 변경하면 안 되는 확정 source action"],
        ["postfix", "prefix 뒤에 새로 생성되어 실행 큐에 추가되는 action"],
        ["logical tensor", "현재 로봇에 의미 있는 32×16 action 영역"],
        ["model tensor", "GR00T action head 고정 최대 규격 40×132"],
        ["source clock", "모델 waypoint 시간축. 현재 15 Hz"],
        ["control clock", "로봇 명령 발행 시간축. 15 Hz 이상"],
        ["bootstrap", "빈 큐에서 d=0으로 수행하는 최초 요청"],
        ["Actor", "관측에서 행동 또는 action chunk를 생성하는 정책. VLA가 이 역할을 수행할 수 있음"],
        ["Critic", "상태 또는 상태-행동의 장기 누적 보상을 추정하는 V/Q 함수"],
        ["Advantage", "선택한 행동이 현재 정책의 평균 행동보다 얼마나 유리했는지 나타내는 값"],
        ["Reward / Cost", "달성할 성능 신호 / 넘지 않아야 할 안전·위험 신호"],
    ]
    add_table(doc, glossary, "표 2 — 용어와 표기")
    doc.add_page_break()

    doc.add_paragraph("작성 근거와 문서 양식", style="Heading 1")
    p = doc.add_paragraph()
    add_inline_text(
        p,
        "본문은 공개된 Dara’s Technical Project Report Template의 Cover–Abstract–Contents–Background–Analysis–Design–Implementation–Testing–Outcome–Conclusion–References–Appendices 구조와 Word style을 기반으로 작성했다. 한국어 가독성을 위해 글꼴과 일부 표 스타일만 현지화했다.",
    )
    p = doc.add_paragraph()
    add_inline_text(
        p,
        "장 구성과 번호 체계는 University of Melbourne 기술보고서 지침 및 MIT Technical Report Format을 교차 참고했다. 기술 사실은 공개 형식보다 로컬 checkpoint와 source code를 우선했다.",
    )
    doc.add_page_break()


def add_body_from_markdown(doc: Document, lines: list[str]) -> None:
    start = next(i for i, line in enumerate(lines) if line.startswith("# 1. 서론"))
    lines = lines[start:]
    index = 0
    table_no = 2
    last_heading = "기술 데이터"
    pending_figure_caption = False
    in_code = False
    code_lines: list[str] = []
    paragraph_lines: list[str] = []

    def flush_paragraph() -> None:
        if not paragraph_lines:
            return
        text = " ".join(item.strip() for item in paragraph_lines).strip()
        paragraph_lines.clear()
        if not text:
            return
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0.55)
        add_inline_text(p, text)

    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()
        if in_code:
            if stripped.startswith("```"):
                p = doc.add_paragraph(style="Code")
                p.paragraph_format.left_indent = Cm(0.5)
                p.paragraph_format.right_indent = Cm(0.3)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(7)
                p.add_run("\n".join(code_lines))
                code_lines.clear()
                in_code = False
            else:
                code_lines.append(raw)
            index += 1
            continue
        if stripped.startswith("```"):
            flush_paragraph()
            in_code = True
            index += 1
            continue
        if not stripped:
            flush_paragraph()
            index += 1
            continue
        if stripped.startswith("# "):
            flush_paragraph()
            p = doc.add_paragraph(stripped[2:], style="Heading 1")
            p.paragraph_format.page_break_before = True
            last_heading = stripped[2:]
            index += 1
            continue
        if stripped.startswith("## "):
            flush_paragraph()
            last_heading = stripped[3:]
            doc.add_paragraph(last_heading, style="Heading 2")
            index += 1
            continue
        if stripped.startswith("### "):
            flush_paragraph()
            last_heading = stripped[4:]
            doc.add_paragraph(last_heading, style="Heading 3")
            index += 1
            continue
        image_match = re.fullmatch(r"!\[(.*?)\]\((.*?)\)", stripped)
        if image_match:
            flush_paragraph()
            path = resolve_image(image_match.group(2))
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            width = Cm(16.2) if "diagram" in path.name else Cm(15.0)
            p.add_run().add_picture(str(path), width=width)
            pending_figure_caption = True
            index += 1
            continue
        if pending_figure_caption and stripped.startswith("그림 "):
            add_caption(doc, stripped)
            pending_figure_caption = False
            index += 1
            continue
        if stripped == "$$":
            flush_paragraph()
            equation_lines: list[str] = []
            index += 1
            while index < len(lines) and lines[index].strip() != "$$":
                equation_lines.append(lines[index].strip())
                index += 1
            if index >= len(lines):
                raise ValueError("닫히지 않은 Markdown 수식 블록이 있습니다.")
            add_equation(doc, "$$\n" + "\n".join(equation_lines) + "\n$$")
            index += 1
            continue
        if stripped.startswith("$$") and stripped.endswith("$$"):
            flush_paragraph()
            add_equation(doc, stripped)
            index += 1
            continue
        if stripped.startswith("|"):
            flush_paragraph()
            rows, index = parse_table(lines, index)
            table_no += 1
            add_table(doc, rows, f"표 {table_no} — {last_heading}")
            continue
        if re.match(r"^[-*] ", stripped):
            flush_paragraph()
            p = doc.add_paragraph(style="List Bullet") if "List Bullet" in doc.styles else doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.65)
            p.paragraph_format.first_line_indent = Cm(-0.35)
            add_inline_text(p, "• " + stripped[2:])
            index += 1
            continue
        if re.match(r"^\d+\. ", stripped):
            flush_paragraph()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.65)
            p.paragraph_format.first_line_indent = Cm(-0.45)
            add_inline_text(p, stripped)
            index += 1
            continue
        if stripped.startswith("> "):
            flush_paragraph()
            p = doc.add_paragraph(style="Quote")
            add_inline_text(p, stripped[2:])
            index += 1
            continue
        paragraph_lines.append(stripped)
        index += 1
    flush_paragraph()


def add_headers_footers(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(2.1)
        section.right_margin = Cm(2.0)
        section.header_distance = Cm(0.8)
        section.footer_distance = Cm(0.8)
        header = section.header.paragraphs[0]
        header.text = "DRILL VLA · TT-RTC · RL TECHNICAL REPORT"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if header.runs:
            header.runs[0].font.size = Pt(7.5)
            header.runs[0].font.color.rgb = RGBColor.from_string(COLORS["gray"])
        footer = section.footer.paragraphs[0]
        footer.clear()
        add_page_number(footer)


def set_update_fields(doc: Document) -> None:
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def build() -> Path:
    make_diagrams()
    convert_template_package()
    doc = Document(WORKING_TEMPLATE)
    clear_document(doc)
    configure_styles(doc)
    doc.core_properties.title = "Drill 삽입 작업을 위한 VLA·TT-RTC·강화학습 기술보고서"
    doc.core_properties.subject = "VLA, TT-RTC, reinforcement learning, loss, runtime, and validation"
    doc.core_properties.author = "CYCLO Intelligence Engineering"
    doc.core_properties.keywords = "TT-RTC, VLA, Drill, reinforcement learning, PPO, SAC, offline RL"
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    add_cover(doc)
    add_front_matter(doc, lines)
    add_body_from_markdown(doc, lines)
    add_headers_footers(doc)
    set_update_fields(doc)
    doc.save(OUTPUT_DOCX)
    return OUTPUT_DOCX


if __name__ == "__main__":
    print(build())
